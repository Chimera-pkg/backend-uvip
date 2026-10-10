import os
import uuid
from datetime import datetime
from typing import List, Optional
from uuid import UUID
import json
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form, Query
from fastapi.responses import JSONResponse, StreamingResponse
from sqlalchemy.orm import Session
from geoalchemy2 import WKTElement

from app.db.database import get_db
from app.db.models import (
    StreetPhoto,
    SegmentationResult,
    PerceptionPrediction,
    ShapValue,
    SimulationSession,
    SimulationResult,
    PolicyRecommendation,
    OfflineSyncQueue,
    Project,
    User,
    VideoOutputSegmentation
)
from app.db.enums import PhotoSource, ProcessingStatus
from app.routers.auth import get_current_user
from app.routers.segmentation_results import create_segmentation
from app.schemas.street_photo import StreetPhotoResponse, StreetPhotoUpdate, PaginatedStreetPhotoResponse
from app.schemas.segmentation_result import SegmentationResultCreate

# pagination
from math import ceil

# background task
from fastapi import BackgroundTasks
from app.service.ai_service import process_photo_with_ai_task

# export excel
import io
import tempfile
import logging
from pathlib import Path

import httpx
from PIL import Image as PILImage
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
from fastapi import APIRouter, Depends, HTTPException, Query
# from fastapi.responses import StreamingResponse

router = APIRouter(prefix="/street-photos", tags=["Street Photos"])

# Folder tujuan penyimpanan foto di server
UPLOAD_DIR = "uploads/photos"
os.makedirs(UPLOAD_DIR, exist_ok=True)

@router.post("/", response_model=StreetPhotoResponse, status_code=status.HTTP_201_CREATED)
async def upload_street_photo(
    # Injeksi Background Tasks
    background_tasks: BackgroundTasks,
    
    # File Fisik dari Client
    file: UploadFile = File(...),
    
    source: PhotoSource = Form(...),
    latitude: float = Form(...),
    longitude: float = Form(...),
    street_name: Optional[str] = Form(None),
    captured_at: datetime = Form(...),
    project_id: UUID = Form(...),
    mission_id: Optional[UUID] = Form(None),
    gps_accuracy_m: Optional[float] = Form(None),
    compass_azimuth: Optional[float] = Form(None),
    exif_timestamp: Optional[datetime] = Form(None),
    is_manual_capture: bool = Form(False),
    is_offline_sync: bool = Form(False),
    
    # Injeksi DB & User
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    # Validation Format File
    allowed_extensions = {".jpg", ".jpeg", ".png", ".webp"}
    file_ext = os.path.splitext(file.filename)[1].lower()
    if file_ext not in allowed_extensions:
        raise HTTPException(
            status_code=400, 
            detail="Format file tidak didukung! Gunakan .jpg, .jpeg, .png, atau .webp"
        )

    # Buat nama file unik menggunakan UUID
    saved_filename = f"{uuid.uuid4()}{file_ext}"
    relative_file_path = os.path.join(UPLOAD_DIR, saved_filename).replace("\\", "/")

    # Simpan file ke direktori server & hitung ukurannya
    try:
        contents = await file.read()
        file_size_kb = int(len(contents) / 1024) 
        with open(relative_file_path, "wb") as f:
            f.write(contents)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gagal menyimpan file: {str(e)}")

    # Buat titik geometri PostGIS (Point SRID 4326: Longitude dulu baru Latitude!)
    geom_point = WKTElement(f"POINT({longitude} {latitude})", srid=4326)

    # Validasi project_id jika dikirim
    if project_id:
        if not db.query(Project).filter(Project.id == project_id).first():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Project dengan ID '{project_id}' tidak ditemukan."
            )

    # Simpan record ke database
    photo = StreetPhoto(
        project_id=project_id,
        mission_id=mission_id,
        uploaded_by=current_user.id,
        source=source,
        original_filename=file.filename,
        file_path=relative_file_path,
        file_size_kb=file_size_kb,
        latitude=latitude,
        longitude=longitude,
        street_name=street_name,
        geom=geom_point,
        gps_accuracy_m=gps_accuracy_m,
        compass_azimuth=compass_azimuth,
        exif_timestamp=exif_timestamp,
        is_manual_capture=is_manual_capture,
        is_offline_sync=is_offline_sync,
        captured_at=captured_at,
        processing_status=ProcessingStatus.QUEUED # Status awal antrean
    )

    db.add(photo)
    db.commit()
    db.refresh(photo)

    # Lempar eksekusi pemanggilan API AI ke background thread
    background_tasks.add_task(process_photo_with_ai_task, photo.id, relative_file_path)

    return photo

@router.post("/bulk", response_model=list[StreetPhotoResponse], status_code=status.HTTP_201_CREATED)
async def upload_multiple_street_photos(
    background_tasks: BackgroundTasks,
    
    # 1. Menerima banyak file foto sekaligus
    files: List[UploadFile] = File(..., description="Daftar file foto"),
    
    # 2. Menerima metadata dalam bentuk JSON string (berisi array objek data per foto)
    # Contoh format JSON: [{"latitude": -7.2, "longitude": 112.6, "captured_at": "2026-10-06T10:00:00Z", "source": "MOBILE", ...}, {...}]
    metadata_json: str = Form(..., description="JSON string berisi list metadata yang urutannya sesuai dengan files"),
    
    project_id: UUID = Form(..., description="Project ID yang berlaku untuk semua foto batch ini (atau bisa dipindah ke JSON jika beda-beda)"),
    
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    # Validasi Project ID
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Project dengan ID '{project_id}' tidak ditemukan."
        )

    # Parse JSON metadata
    try:
        metadata_list = json.loads(metadata_json)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="Format metadata_json tidak valid (bukan JSON yang benar).")

    # Validasi jumlah file dan metadata harus klop/sama banyak
    if len(files) != len(metadata_list):
        raise HTTPException(
            status_code=400, 
            detail=f"Jumlah file ({len(files)}) tidak sinkron dengan jumlah data metadata ({len(metadata_list)})."
        )

    allowed_extensions = {".jpg", ".jpeg", ".png", ".webp"}
    uploaded_photos = []

    # Looping berdasarkan indeks supaya file dan metadatanya berpasangan dengan pas
    for index, file in enumerate(files):
        meta = metadata_list[index]
        
        file_ext = os.path.splitext(file.filename)[1].lower() if file.filename else ""
        if file_ext not in allowed_extensions:
            raise HTTPException(
                status_code=400, 
                detail=f"Format file '{file.filename}' tidak didukung!"
            )

        saved_filename = f"{uuid.uuid4()}{file_ext}"
        relative_file_path = os.path.join(UPLOAD_DIR, saved_filename).replace("\\", "/")

        # Simpan file fisik
        try:
            contents = await file.read()
            file_size_kb = int(len(contents) / 1024) 
            with open(relative_file_path, "wb") as f:
                f.write(contents)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Gagal menyimpan file {file.filename}: {str(e)}")
        finally:
            await file.close()

        # Ambil koordinat unik dari metadata spesifik foto ini
        lat = meta.get("latitude")
        lon = meta.get("longitude")
        if lat is None or lon is None:
            raise HTTPException(status_code=400, detail=f"Latitude atau Longitude kosong pada indeks ke-{index}")

        geom_point = WKTElement(f"POINT({lon} {lat})", srid=4326)

        # Simpan ke database menggunakan data spesifik dari `meta`
        photo = StreetPhoto(
            project_id=project_id,
            mission_id=meta.get("mission_id"),
            uploaded_by=current_user.id,
            source=meta.get("source", PhotoSource.MOBILE_LIVE),
            original_filename=file.filename,
            file_path=relative_file_path,
            file_size_kb=file_size_kb,
            latitude=lat,
            longitude=lon,
            street_name=meta.get("street_name"),
            geom=geom_point,
            gps_accuracy_m=meta.get("gps_accuracy_m"),
            compass_azimuth=meta.get("compass_azimuth"),
            exif_timestamp=meta.get("exif_timestamp"),
            is_manual_capture=meta.get("is_manual_capture", False),
            is_offline_sync=meta.get("is_offline_sync", False),
            captured_at=meta.get("captured_at"),
            processing_status=ProcessingStatus.QUEUED
        )

        db.add(photo)
        db.commit()
        db.refresh(photo)

        # Daftarkan background task AI secara independen
        background_tasks.add_task(process_photo_with_ai_task, photo.id, relative_file_path)
        uploaded_photos.append(photo)

    return uploaded_photos

# 2. READ ALL
@router.get("/", response_model=PaginatedStreetPhotoResponse)
def list_photos(
    page: int = Query(1, ge=1, description="Nomor halaman yang ingin diakses"),
    size: int = Query(10, ge=1, le=100, description="Jumlah data maksimal per halaman"),
    project_id: Optional[str] = Query(None, description="Filter berdasarkan project (UUID atau 'unassigned' untuk media tanpa project)"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    base_query = db.query(StreetPhoto).filter(
        StreetPhoto.file_path.ilike(f"{UPLOAD_DIR}/%")
    )

    if project_id:
        if project_id.lower() == "unassigned":
            base_query = base_query.filter(StreetPhoto.project_id.is_(None))
        else:
            base_query = base_query.filter(StreetPhoto.project_id == project_id)

    total_data = base_query.count()
    total_pages = ceil(total_data / size) if total_data > 0 else 1
    skip = (page - 1) * size
    photos = base_query.order_by(StreetPhoto.created_at.desc()).offset(skip).limit(size).all()

    return {
        "total_data": total_data,
        "total_pages": total_pages,
        "current_page": page,
        "data": photos
    }


# 3. READ BY ID
@router.get("/{photo_id}", response_model=StreetPhotoResponse)
def get_photo(photo_id: UUID, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    photo = db.query(StreetPhoto).filter(StreetPhoto.id == photo_id).first()
    if not photo:
        raise HTTPException(status_code=404, detail="Foto tidak ditemukan")
    return photo


# 4. UPDATE BY ID
@router.put("/{photo_id}", response_model=StreetPhotoResponse)
def update_photo(photo_id: UUID, data: StreetPhotoUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    photo = db.query(StreetPhoto).filter(StreetPhoto.id == photo_id).first()
    if not photo:
        raise HTTPException(status_code=404, detail="Foto tidak ditemukan")
    
    update_data = data.model_dump(exclude_unset=True)
    
    # Update otomatis geom jika latitude/longitude berubah
    lat = update_data.get("latitude", photo.latitude)
    long = update_data.get("longitude", photo.longitude)
    if "latitude" in update_data or "longitude" in update_data:
        update_data["geom"] = WKTElement(f"POINT({long} {lat})", srid=4326)

    for key, val in update_data.items():
        setattr(photo, key, val)

    db.commit()
    db.refresh(photo)
    return photo


# 5. DELETE BY ID (TERMASUK HAPUS FILE FISIK)
@router.delete("/{photo_id}")
def delete_photo(
    photo_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    execute_full_cascade_delete_photo(photo_id=photo_id, db=db)
    
    return JSONResponse(
        status_code=status.HTTP_200_OK,
        content={
            "status": "success",
            "message": "File foto beserta seluruh riwayat analisis & simulasinya berhasil dihapus bersih",
            "deleted_id": str(photo_id)
        }
    )

def execute_full_cascade_delete_photo(photo_id: UUID, db: Session) -> None:
    """
    Menghapus record StreetPhoto beserta seluruh dependensi tabel anak
    secara berurutan dari level terdalam untuk menghindari ForeignKeyViolation.
    """
    photo = db.query(StreetPhoto).filter(StreetPhoto.id == photo_id).first()
    if not photo:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Foto/Video dengan ID '{photo_id}' tidak ditemukan."
        )

    db.query(OfflineSyncQueue).filter(
        OfflineSyncQueue.synced_photo_id == photo_id
    ).update({"synced_photo_id": None}, synchronize_session=False)

    predictions = db.query(PerceptionPrediction).filter(
        PerceptionPrediction.photo_id == photo_id
    ).all()
    prediction_ids = [p.id for p in predictions]

    sim_query = db.query(SimulationSession).filter(
        (SimulationSession.base_photo_id == photo_id) |
        (SimulationSession.base_prediction_id.in_(prediction_ids) if prediction_ids else False)
    )
    sim_sessions = sim_query.all()
    session_ids = [s.id for s in sim_sessions]

    # delete anak simulation_session
    if session_ids:
        # delete policy_recommendations
        db.query(PolicyRecommendation).filter(
            PolicyRecommendation.session_id.in_(session_ids)
        ).delete(synchronize_session=False)

        # delete simulation_results
        db.query(SimulationResult).filter(
            SimulationResult.session_id.in_(session_ids)
        ).delete(synchronize_session=False)

        # delete simulation_sessions
        db.query(SimulationSession).filter(
            SimulationSession.id.in_(session_ids)
        ).delete(synchronize_session=False)

    # delete perception_predictions
    if prediction_ids:
        db.query(ShapValue).filter(
            ShapValue.prediction_id.in_(prediction_ids)
        ).delete(synchronize_session=False)

        db.query(PerceptionPrediction).filter(
            PerceptionPrediction.id.in_(prediction_ids)
        ).delete(synchronize_session=False)

    # delete segmentation
    segmentation = db.query(SegmentationResult).filter(
        SegmentationResult.photo_id == photo_id
    ).first()

    if segmentation:
        if segmentation.mask_file_path and os.path.exists(segmentation.mask_file_path):
            try:
                os.remove(segmentation.mask_file_path)
            except OSError:
                pass
        db.delete(segmentation)

    # delete file fisik
    if photo.file_path and os.path.exists(photo.file_path):
        try:
            os.remove(photo.file_path)
        except OSError:
            pass

    # delete video_output_segmentations
    db.query(VideoOutputSegmentation).filter(
            VideoOutputSegmentation.photo_id == photo_id
        ).delete(synchronize_session=False)
        # ----------------------------------------------------

    # delete file fisik
    if photo.file_path and os.path.exists(photo.file_path):
        try:
            os.remove(photo.file_path)
        except OSError:
            pass

    db.delete(photo)
    db.commit()

# export excel
logger = logging.getLogger(__name__)

# router = APIRouter(prefix="/street-photos", tags=["Street Photos"])

BASE_DIR = Path(__file__).resolve().parent.parent.parent
UPLOAD_DIR = "uploads/photos"
# LOCAL_BASE_URL = "http://127.0.0.1:8000"
LOCAL_BASE_URL = "http://80.241.214.39"
AI_BASE_URL = "https://trends-womens-catherine-approval.trycloudflare.com"
IMG_W, IMG_H = 90, 90


def _load_local_image(file_path: str) -> Optional[PILImage.Image]:
    rel = file_path.lstrip("/\\").replace("\\", "/")
    abs_path = BASE_DIR / rel
    if not abs_path.exists():
        logger.warning(f"[EXPORT] File lokal tidak ditemukan: {abs_path}")
        return None
    try:
        img = PILImage.open(abs_path)
        img.load()
        return img
    except Exception as e:
        logger.exception(f"[EXPORT] Gagal buka {abs_path}: {e}")
        return None


def _load_remote_image(url: str, timeout: float = 15.0) -> Optional[PILImage.Image]:
    try:
        with httpx.Client(timeout=timeout, follow_redirects=True) as client:
            resp = client.get(url)
        if resp.status_code != 200:
            logger.warning(f"[EXPORT] HTTP {resp.status_code} untuk {url}")
            return None
        img = PILImage.open(io.BytesIO(resp.content))
        img.load()
        return img
    except Exception as e:
        logger.exception(f"[EXPORT] Gagal download {url}: {e}")
        return None


def _save_temp_png(pil_img: PILImage.Image) -> str:
    """
    Normalisasi PIL Image, resize, dan simpan sebagai file PNG di temp.
    Return: path file temp.
    """
    # Normalisasi mode
    if pil_img.mode in ("RGBA", "LA", "P"):
        bg = PILImage.new("RGB", pil_img.size, (255, 255, 255))
        if pil_img.mode == "P":
            pil_img = pil_img.convert("RGBA")
        bg.paste(
            pil_img,
            mask=pil_img.split()[-1] if pil_img.mode in ("RGBA", "LA") else None,
        )
        pil_img = bg
    elif pil_img.mode != "RGB":
        pil_img = pil_img.convert("RGB")

    # Resize
    pil_img = pil_img.resize((IMG_W * 2, IMG_H * 2))

    # Simpan ke temp file
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    pil_img.save(tmp.name, format="PNG")
    tmp.close()
    return tmp.name


# @router.get(
#     "/export/export-photos-excel",
#     summary="",
# )
# def export_street_photos_excel(
#     project_id: Optional[UUID] = Query(None),
#     db: Session = Depends(get_db),
#     current_user: User = Depends(get_current_user),
# ):
#     base_query = db.query(StreetPhoto).filter(
#         StreetPhoto.file_path.ilike(f"{UPLOAD_DIR}/%")
#     )
#     if project_id:
#         base_query = base_query.filter(StreetPhoto.project_id == project_id)

#     photos = base_query.order_by(StreetPhoto.created_at.desc()).all()
#     if not photos:
#         raise HTTPException(status_code=404, detail="Tidak ada data foto untuk diexport.")

#     wb = Workbook()
#     ws = wb.active
#     ws.title = "Laporan Foto & Analisis"

#     headers = [
#         "No", "Foto Asli", "Foto Overlay AI", "Latitude", "Longitude", "Waktu Tangkap",
#         "Vegetasi (%)", "Bangunan (%)", "Jalan (%)", "Langit (%)",
#         "Walkability Ratio", "Visual Clutter", "Beauty Score",
#         "Safety Score", "Comfort Score", "UVI Score",
#     ]
#     ws.append(headers)

#     widths = {1: 6, 2: 18, 3: 18, 4: 15, 5: 15, 6: 22}
#     for col, w in widths.items():
#         ws.column_dimensions[get_column_letter(col)].width = w

#     # Simpan path temp untuk cleanup nanti
#     temp_files: List[str] = []

#     try:
#         for index, raw_photo in enumerate(photos, start=2):
#             segmentation = db.query(SegmentationResult).filter(
#                 SegmentationResult.photo_id == raw_photo.id
#             ).first()
#             prediction = db.query(PerceptionPrediction).filter(
#                 PerceptionPrediction.photo_id == raw_photo.id
#             ).first()

#             ws.row_dimensions[index].height = IMG_H * 0.75 + 5

#             ws.cell(row=index, column=1, value=index - 1)
#             ws.cell(row=index, column=4, value=raw_photo.latitude)
#             ws.cell(row=index, column=5, value=raw_photo.longitude)
#             ws.cell(row=index, column=6, value=str(raw_photo.captured_at) if raw_photo.captured_at else "")

#             ws.cell(row=index, column=7, value=segmentation.vegetation_pct if segmentation else None)
#             ws.cell(row=index, column=8, value=segmentation.building_pct if segmentation else None)
#             ws.cell(row=index, column=9, value=segmentation.road_pct if segmentation else None)
#             ws.cell(row=index, column=10, value=segmentation.sky_pct if segmentation else None)
#             ws.cell(row=index, column=11, value=segmentation.walkability_ratio if segmentation else None)
#             ws.cell(row=index, column=12, value=segmentation.visual_clutter_index if segmentation else None)
#             ws.cell(row=index, column=13, value=prediction.beauty_score if prediction else None)
#             ws.cell(row=index, column=14, value=prediction.safety_score if prediction else None)
#             ws.cell(row=index, column=15, value=prediction.comfort_score if prediction else None)
#             ws.cell(row=index, column=16, value=prediction.uvi_score if prediction else None)

#             # === FOTO ASLI ===
#             if raw_photo.file_path:
#                 pil_img = _load_local_image(raw_photo.file_path)
#                 if pil_img:
#                     try:
#                         tmp_path = _save_temp_png(pil_img)
#                         temp_files.append(tmp_path)
#                         xl_img = OpenpyxlImage(tmp_path)
#                         xl_img.width = IMG_W
#                         xl_img.height = IMG_H
#                         ws.add_image(xl_img, f"B{index}")
#                     except Exception as e:
#                         logger.exception(f"[EXPORT] add_image lokal gagal: {e}")
#                         ws.cell(row=index, column=2, value="Gagal render")
#                 else:
#                     rel = raw_photo.file_path.lstrip("/\\").replace("\\", "/")
#                     ws.cell(row=index, column=2, value=f"{LOCAL_BASE_URL}/{rel}")
#             else:
#                 ws.cell(row=index, column=2, value="Path kosong")

#             # === FOTO OVERLAY AI ===
#             if segmentation and segmentation.segmentation_overlay_url:
#                 clean_path = segmentation.segmentation_overlay_url.lstrip("/")
#                 overlay_url = f"{AI_BASE_URL}/{clean_path}"
#                 pil_overlay = _load_remote_image(overlay_url)
#                 if pil_overlay:
#                     try:
#                         tmp_path = _save_temp_png(pil_overlay)
#                         temp_files.append(tmp_path)
#                         xl_img = OpenpyxlImage(tmp_path)
#                         xl_img.width = IMG_W
#                         xl_img.height = IMG_H
#                         ws.add_image(xl_img, f"C{index}")
#                     except Exception as e:
#                         logger.exception(f"[EXPORT] add_image overlay gagal: {e}")
#                         ws.cell(row=index, column=3, value="Gagal render")
#                 else:
#                     ws.cell(row=index, column=3, value="Gagal ambil overlay")
#             else:
#                 ws.cell(row=index, column=3, value="Belum ada overlay")

#         output = io.BytesIO()
#         wb.save(output)
#         output.seek(0)

#     finally:
#         # Cleanup semua file temp
#         for p in temp_files:
#             try:
#                 os.remove(p)
#             except OSError:
#                 pass

#     filename = "Laporan_Street_Photos_Lengkap.xlsx"
#     return StreamingResponse(
#         output,
#         media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
#         headers={"Content-Disposition": f'attachment; filename="{filename}"'},
#     )
@router.get(
    "/export/export-photos-excel",
    summary="",
)
def export_street_photos_excel(
    project_id: Optional[UUID] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    base_query = db.query(StreetPhoto).filter(
        StreetPhoto.file_path.ilike(f"{UPLOAD_DIR}/%")
    )
    if project_id:
        base_query = base_query.filter(StreetPhoto.project_id == project_id)

    # ← ORDER BY FILENAME
    photos = base_query.order_by(StreetPhoto.original_filename.asc()).all()
    if not photos:
        raise HTTPException(status_code=404, detail="Tidak ada data foto untuk diexport.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Foto & Analisis"

    headers = [
        "No", "Foto Asli", "Foto Overlay AI", "Nama Foto",
        "Latitude", "Longitude", "Waktu Tangkap",
        "Vegetasi (%)", "Bangunan (%)", "Jalan (%)", "Langit (%)",
        "Walkability Ratio", "Visual Clutter", "Beauty Score",
        "Safety Score", "Comfort Score", "UVI Score",
    ]
    ws.append(headers)

    # ← TAMBAH KOLOM D
    widths = {1: 6, 2: 18, 3: 18, 4: 25, 5: 15, 6: 15, 7: 22}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    temp_files: List[str] = []

    try:
        for index, raw_photo in enumerate(photos, start=2):
            segmentation = db.query(SegmentationResult).filter(
                SegmentationResult.photo_id == raw_photo.id
            ).first()
            prediction = db.query(PerceptionPrediction).filter(
                PerceptionPrediction.photo_id == raw_photo.id
            ).first()

            ws.row_dimensions[index].height = IMG_H * 0.75 + 5

            ws.cell(row=index, column=1, value=index - 1)
            # ← KOLOM BARU: NAMA FOTO
            ws.cell(row=index, column=4, value=raw_photo.original_filename)
            ws.cell(row=index, column=5, value=raw_photo.latitude)
            ws.cell(row=index, column=6, value=raw_photo.longitude)
            ws.cell(row=index, column=7, value=str(raw_photo.captured_at) if raw_photo.captured_at else "")

            ws.cell(row=index, column=8, value=segmentation.vegetation_pct if segmentation else None)
            ws.cell(row=index, column=9, value=segmentation.building_pct if segmentation else None)
            ws.cell(row=index, column=10, value=segmentation.road_pct if segmentation else None)
            ws.cell(row=index, column=11, value=segmentation.sky_pct if segmentation else None)
            ws.cell(row=index, column=12, value=segmentation.walkability_ratio if segmentation else None)
            ws.cell(row=index, column=13, value=segmentation.visual_clutter_index if segmentation else None)
            ws.cell(row=index, column=14, value=prediction.beauty_score if prediction else None)
            ws.cell(row=index, column=15, value=prediction.safety_score if prediction else None)
            ws.cell(row=index, column=16, value=prediction.comfort_score if prediction else None)
            ws.cell(row=index, column=17, value=prediction.uvi_score if prediction else None)

            # === FOTO ASLI ===
            if raw_photo.file_path:
                pil_img = _load_local_image(raw_photo.file_path)
                if pil_img:
                    try:
                        tmp_path = _save_temp_png(pil_img)
                        temp_files.append(tmp_path)
                        xl_img = OpenpyxlImage(tmp_path)
                        xl_img.width = IMG_W
                        xl_img.height = IMG_H
                        ws.add_image(xl_img, f"B{index}")
                    except Exception as e:
                        logger.exception(f"[EXPORT] add_image lokal gagal: {e}")
                        ws.cell(row=index, column=2, value="Gagal render")
                else:
                    rel = raw_photo.file_path.lstrip("/\\").replace("\\", "/")
                    ws.cell(row=index, column=2, value=f"{LOCAL_BASE_URL}/{rel}")
            else:
                ws.cell(row=index, column=2, value="Path kosong")

            # === FOTO OVERLAY AI ===
            if segmentation and segmentation.segmentation_overlay_url:
                clean_path = segmentation.segmentation_overlay_url.lstrip("/")
                overlay_url = f"{AI_BASE_URL}/{clean_path}"
                pil_overlay = _load_remote_image(overlay_url)
                if pil_overlay:
                    try:
                        tmp_path = _save_temp_png(pil_overlay)
                        temp_files.append(tmp_path)
                        xl_img = OpenpyxlImage(tmp_path)
                        xl_img.width = IMG_W
                        xl_img.height = IMG_H
                        ws.add_image(xl_img, f"C{index}")
                    except Exception as e:
                        logger.exception(f"[EXPORT] add_image overlay gagal: {e}")
                        ws.cell(row=index, column=3, value="Gagal render")
                else:
                    ws.cell(row=index, column=3, value="Gagal ambil overlay")
            else:
                ws.cell(row=index, column=3, value="Belum ada overlay")

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

    finally:
        for p in temp_files:
            try:
                os.remove(p)
            except OSError:
                pass

    filename = "Laporan_Street_Photos_Lengkap.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )